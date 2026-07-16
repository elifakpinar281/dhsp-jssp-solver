#!/usr/bin/env python3

import io
import json
import re
import sys
import zipfile
import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else "Demoanlage.xlsx"
N_CARRIERS = int(sys.argv[2]) if len(sys.argv) > 2 else 10

COL_BAD = 2
COL_DESC = 3
TYPE_MIN_COLS = [4, 6, 8, 10, 12]
TYPE_MAX_COLS = [5, 7, 9, 11, 13]
TYPE_COUNT = 5

FALLBACK_PROPORTIONS = [0.4, 0.2, 0.1, 0.2, 0.1]


def load_workbook_robust(path):
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except TypeError:
        src = zipfile.ZipFile(path)
        buf = io.BytesIO()
        dst = zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED)
        for name in src.namelist():
            content = src.read(name)
            if name == "xl/styles.xml":
                text = content.decode("utf-8")
                text = re.sub(r"<extLst>.*?</extLst>", "", text, flags=re.DOTALL)
                content = text.encode("utf-8")
            dst.writestr(name, content)
        dst.close()
        buf.seek(0)
        return openpyxl.load_workbook(buf, data_only=True)


def load_rows():
    wb = load_workbook_robust(XLSX)
    ws = wb["Fahrplanmix"]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(r)
    return rows


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def bath_count(bad):
    m = re.match(r"^\s*(\d+)\s*-\s*(\d+)\s*$", str(bad))
    if not m:
        return 1
    lo, hi = int(m.group(1)), int(m.group(2))
    return (hi - lo) // 10 + 1


def parse_max(v):
    if is_number(v):
        return int(v)
    return "inf"


def find_proportions(rows):
    for r in rows:
        values = []
        for c in TYPE_MIN_COLS:
            v = r[c] if len(r) > c else None
            if not is_number(v):
                values = []
                break
            values.append(float(v))
        if len(values) != TYPE_COUNT:
            continue

        total = sum(values)
        if abs(total - 1.0) > 0.001:
            continue

        for v in values:
            if v <= 0.0 or v >= 1.0:
                values = []
                break
        if len(values) == TYPE_COUNT:
            return values, True

    return list(FALLBACK_PROPORTIONS), False


def split_carriers(proportions, n_carriers):
    counts = []
    remainders = []
    for t in range(TYPE_COUNT):
        exact = proportions[t] * n_carriers
        base = int(exact)
        counts.append(base)
        remainders.append((exact - base, t))

    missing = n_carriers - sum(counts)
    remainders.sort(key=lambda pair: pair[0], reverse=True)
    for i in range(missing):
        t = remainders[i % TYPE_COUNT][1]
        counts[t] = counts[t] + 1
    return counts


def build_task_rows(rows):
    task_rows = []
    for r in rows:
        bad = r[COL_BAD] if len(r) > COL_BAD else None
        desc = r[COL_DESC] if len(r) > COL_DESC else None
        if bad is None:
            continue
        if not (is_number(bad) or (isinstance(bad, str) and re.search(r"\d", bad))):
            continue
        mins = []
        maxs = []
        for c in TYPE_MIN_COLS:
            v = r[c] if len(r) > c else None
            mins.append(v if is_number(v) else None)
        for c in TYPE_MAX_COLS:
            v = r[c] if len(r) > c else None
            maxs.append(parse_max(v))
        if any(m is not None for m in mins):
            task_rows.append((str(bad).strip(), str(desc).strip(), mins, maxs))
    return task_rows


def main():
    rows = load_rows()
    proportions, from_excel = find_proportions(rows)
    task_rows = build_task_rows(rows)

    machine_id = {}
    machine_name = {}
    capacities = []
    for bad, desc, _, _ in task_rows:
        key = bad
        if key not in machine_id:
            idx = len(machine_id)
            machine_id[key] = idx
            machine_name[idx] = f"{bad} {desc}"
            capacities.append(bath_count(bad))

    jobtype_ops = {t: [] for t in range(TYPE_COUNT)}
    for bad, desc, mins, maxs in task_rows:
        for t in range(TYPE_COUNT):
            m = mins[t]
            if m is None:
                continue
            if m <= 0:
                continue
            jobtype_ops[t].append((machine_id[bad], int(m), maxs[t]))

    counts = split_carriers(proportions, N_CARRIERS)

    jobs = []
    for t in range(TYPE_COUNT):
        for _ in range(counts[t]):
            jobs.append((t, list(jobtype_ops[t])))

    machine_count = len(machine_id)

    lines = []
    lines.append(f"JOBS {len(jobs)}")
    lines.append(f"MACHINES {machine_count}")
    lines.append("CAPACITIES " + " ".join(str(c) for c in capacities))
    lines.append("BLOCKING true")
    for _, ops in jobs:
        parts = [str(len(ops))]
        for m, d, mx in ops:
            parts.append(str(m))
            parts.append(str(d))
            parts.append(str(mx))
        lines.append(" ".join(parts))
    with open("demoanlage.txt", "w") as f:
        f.write("\n".join(lines) + "\n")

    mapping = {
        "instance": "Demoanlage (IM1)",
        "nCarriers": N_CARRIERS,
        "jobtypeProportions": {f"Typ{t+1}": proportions[t] for t in range(TYPE_COUNT)},
        "proportionsSource": "Fahrplanmix" if from_excel else "fallback",
        "jobtypeCounts": {f"Typ{t+1}": counts[t] for t in range(TYPE_COUNT)},
        "jobtypeNames": {
            "Typ1": "Cu + ChemNi + Tef",
            "Typ2": "ChemNi + Tef",
            "Typ3": "GalvNi + ChemNi + Tef",
            "Typ4": "ChemNi + Tef (kurz)",
            "Typ5": "GalvNi + ChemNi + Tef (lang)",
        },
        "jobTypeOfJob": [t for t, _ in jobs],
        "machineNames": machine_name,
        "machineCount": machine_count,
        "machineCapacities": {machine_name[i]: capacities[i] for i in range(machine_count)},
        "bathCount": sum(capacities),
        "opsPerJob": [len(ops) for _, ops in jobs],
    }
    with open("demoanlage_mapping.json", "w") as f:
        json.dump(mapping, f, indent=2, ensure_ascii=False)

    source = "read from Fahrplanmix" if from_excel else "FALLBACK (percentage row not found in workbook!)"
    print(f"Share per job type ({source}): {proportions}")
    print(f"Task rows (baths visited): {len(task_rows)}")
    print(f"Distinct machines (bath rows): {machine_count}")
    print(f"Baths total (sum of capacities): {sum(capacities)}")
    print(f"Carriers total:  {len(jobs)}  -> {counts} (type 1..5)")
    total_ops = sum(len(ops) for _, ops in jobs)
    print(f"Operations total: {total_ops}")
    print("Ops per job type:", {f"Typ{t+1}": len(jobtype_ops[t]) for t in range(TYPE_COUNT)})
    print("-> wrote demoanlage.txt, demoanlage_mapping.json")


if __name__ == "__main__":
    main()